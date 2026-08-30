"""Slide-deck business logic: streamed upload persistence, pool merge/dedup,
and the lecture-matching heuristic behind /slides/suggest (A2).

Kept out of the route handlers so those stay pure HTTP glue; lazy
`pypdfium2`/`pypdf`/`pptx` imports live here (only paid for once a deck is
actually touched), not in api.py/routes.
"""

import difflib
import hashlib
import json
import logging
import re
import shutil
import uuid
from pathlib import Path

from fastapi import UploadFile

from . import config
from .errors import TooLargeError, ValidationError

log = logging.getLogger("notely.decks")

UPLOAD_CHUNK_SIZE = 1024 * 1024  # 1 MB


def _load_json(path, default):
    try:
        with open(path) as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return default


async def save_upload_stream(
    file: UploadFile, dest: Path, max_bytes: int = config.MAX_UPLOAD_BYTES, allow_empty: bool = False
) -> int:
    """Stream `file` to `dest` in chunks instead of buffering the whole
    upload in RAM (A3). Writes to a `.part` sibling and renames on success,
    so a client abort or an over-cap upload never leaves a partial file at
    `dest`. Raises errors.TooLargeError past max_bytes -- callers map that
    to HTTP 413 via the app-wide NotelyError handler.
    """
    tmp = dest.with_name(dest.name + f".part{uuid.uuid4().hex[:8]}")
    written = 0
    try:
        with open(tmp, "wb") as out:
            while True:
                chunk = await file.read(UPLOAD_CHUNK_SIZE)
                if not chunk:
                    break
                written += len(chunk)
                if written > max_bytes:
                    raise TooLargeError(f"upload exceeds {max_bytes // (1024 * 1024)} MB limit")
                out.write(chunk)
        if written == 0 and not allow_empty:
            raise ValidationError(f"{file.filename or 'upload'}: empty file")
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise
    tmp.replace(dest)
    return written


# --- pool mode: "I don't know which deck covers which lecture" ------------

# merge_pool writes its result back into the pool dir, so the merged deck has
# to be excluded from the merge's own inputs -- otherwise every re-merge
# folds the previous merge into itself. Text-identical pages get deduped
# away, but image-only pages are deliberately never deduped (see merge_pool),
# so those would accumulate one extra copy per run.
MERGED_DECK_NAME = "_merged.pdf"


def _norm_page_text(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").lower()).strip()


def pool_dir_path() -> Path:
    return config.SLIDES_DIR / "_pool"


async def _save_pool_upload(f: UploadFile, dest_dir: Path, max_bytes: int) -> Path:
    name = (f.filename or "").strip()
    if not name.lower().endswith(".pdf"):
        raise ValidationError(f"{name}: pool mode takes PDFs only — export PPTX decks to PDF first")
    dest = dest_dir / Path(name).name
    await save_upload_stream(f, dest, max_bytes)
    return dest


async def save_pool_uploads(files: list[UploadFile], max_bytes: int = config.MAX_UPLOAD_BYTES) -> list[Path]:
    """Stream every uploaded deck into the pool dir, *replacing* whatever was
    there. Pool mode only accepts PDFs (pptx decks would need a LibreOffice
    conversion pass first).

    An upload defines the whole pool rather than adding to it. The pool feeds
    a single merged deck copied to every lecture, so decks left over from a
    previous course would otherwise stay in that merge forever and every
    lecture would silently get a combined deck spanning both courses.

    Uploads land in a staging dir that only replaces the live pool once all
    of them have succeeded -- same temp-then-rename shape as
    save_upload_stream -- so a rejected pptx or an over-cap file leaves the
    existing pool intact instead of destroying it halfway through.
    """
    pool_dir = pool_dir_path()
    pool_dir.parent.mkdir(parents=True, exist_ok=True)
    staging = pool_dir.with_name(f"{pool_dir.name}.new{uuid.uuid4().hex[:8]}")
    staging.mkdir(parents=True)
    try:
        saved = [await _save_pool_upload(f, staging, max_bytes) for f in files]
    except BaseException:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    shutil.rmtree(pool_dir, ignore_errors=True)
    staging.replace(pool_dir)
    return [pool_dir / p.name for p in saved]


def merge_pool(pool_dir: Path | None = None) -> dict:
    """Merge every deck currently in the pool dir, in filename order,
    skipping pages whose normalized text exactly matches one already kept.
    Course decks repeat earlier material heavily (observed: 936 raw pages,
    only 310 unique, on this project's real pool) -- deduping here instead
    of after stage 4's OCR makes the haystack stage 4 actually searches ~3x
    smaller and less ambiguous, for free. Pages with no extractable text
    (scanned/image-only slides) are never deduped against each other --
    collapsing them on an empty-string hash match would wrongly merge
    visually distinct slides, so they're always kept."""
    from pypdf import PdfReader, PdfWriter

    pool_dir = pool_dir or pool_dir_path()
    # MERGED_DECK_NAME lives in this same dir -- never merge it into itself.
    pool_files = sorted(p for p in pool_dir.glob("*.pdf") if p.name != MERGED_DECK_NAME)
    writer = PdfWriter()
    seen_hashes = set()
    scanned_pages = 0
    for p in pool_files:
        reader = PdfReader(p)
        for page in reader.pages:
            scanned_pages += 1
            # pypdf's extract_text() is noisier than pypdfium2's (occasional
            # spurious spaces from kerning) -- fine here since the hash only
            # needs to be *deterministic* per page, not high-fidelity text;
            # pypdfium2 is used instead wherever extracted text is
            # user-visible (scripts/02_extract_slides.py).
            text = _norm_page_text(page.extract_text())
            if text:
                key = hashlib.md5(text.encode()).hexdigest()
                if key in seen_hashes:
                    continue
                seen_hashes.add(key)
            writer.add_page(page)
    total_pages = len(writer.pages)
    merged_path = pool_dir / MERGED_DECK_NAME
    writer.write(str(merged_path))
    return {
        "pool_files": pool_files,
        "merged_path": merged_path,
        "total_pages": total_pages,
        "scanned_pages": scanned_pages,
    }


def distribute_pool_deck(merged_path: Path, lecture_ids) -> None:
    """Same combined deck for every lecture -- stages pair strictly by
    filename, and stage 4's content matching figures out per video which
    slides were actually shown."""
    for lecture_id in lecture_ids:
        shutil.copyfile(merged_path, config.SLIDES_DIR / f"{lecture_id}.pdf")
        (config.SLIDES_DIR / f"{lecture_id}.pptx").unlink(missing_ok=True)


# --- per-lecture upload -----------------------------------------------------


def _record_deck_meta(lecture_id: str, original_name: str | None) -> None:
    meta = _load_json(config.LECTURES_META_PATH, {})
    meta.setdefault(lecture_id, {})["deck_original_name"] = original_name
    with open(config.LECTURES_META_PATH, "w") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)


async def save_deck_for_lectures(
    file: UploadFile, lecture_ids: list[str], max_bytes: int = config.MAX_UPLOAD_BYTES
) -> list[str]:
    """One deck may cover several lectures: stream it once, then copy it
    under each lecture's name."""
    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    if ext not in ("pdf", "pptx"):
        raise ValidationError("only .pdf and .pptx decks are supported")
    config.SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    tmp = config.SLIDES_DIR / f"_upload_{uuid.uuid4().hex}.{ext}"
    await save_upload_stream(file, tmp, max_bytes)
    saved = []
    try:
        for lecture_id in lecture_ids:
            dest = config.SLIDES_DIR / f"{lecture_id}.{ext}"
            shutil.copyfile(tmp, dest)
            # remove a stale deck of the other extension so stage 02 pairs deterministically
            other = config.SLIDES_DIR / f"{lecture_id}.{'pdf' if ext == 'pptx' else 'pptx'}"
            other.unlink(missing_ok=True)
            saved.append(dest.name)
            _record_deck_meta(lecture_id, file.filename)
    finally:
        tmp.unlink(missing_ok=True)
    return saved


# --- lecture-matching suggestion -------------------------------------------


def _norm_words(s: str) -> list[str]:
    return re.sub(r"[^\w\s]", " ", (s or "").lower()).split()


def _extract_preview_text(path: Path, ext: str) -> str:
    """First couple pages/slides of text -- title slide + first content
    slide -- used as a fingerprint for lecture-matching. Unreadable decks
    log a warning and return "" so the UI just falls back to no
    suggestions, rather than the request failing outright."""
    text = ""
    try:
        if ext == "pdf":
            import pypdfium2 as pdfium

            with pdfium.PdfDocument(str(path)) as doc:
                parts = []
                for i in range(min(2, len(doc))):
                    page = doc[i]
                    textpage = page.get_textpage()
                    try:
                        parts.append(textpage.get_text_range())
                    finally:
                        textpage.close()
                        page.close()
                text = " ".join(parts)
        elif ext == "pptx":
            from pptx import Presentation

            prs = Presentation(str(path))
            parts = []
            for slide in list(prs.slides)[:2]:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            text = " ".join(parts)
    except Exception:
        log.warning("could not extract preview text from .%s upload", ext, exc_info=True)
    return text


async def suggest_lectures_for_upload(
    file: UploadFile, max_bytes: int = config.MAX_UPLOAD_BYTES
) -> list[dict]:
    """Suggest which lecture(s) a deck belongs to by comparing its opening
    slides' text against the video titles. A pre-fill for the UI dropdown,
    never a silent decision."""
    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    urls = config.load_video_urls()
    if ext not in ("pdf", "pptx") or not urls:
        return []

    tmp = config.SLIDES_DIR / f"_suggest_{uuid.uuid4().hex}.{ext}"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    try:
        await save_upload_stream(file, tmp, max_bytes, allow_empty=True)
        text = _extract_preview_text(tmp, ext)
    finally:
        tmp.unlink(missing_ok=True)

    deck_words = set(_norm_words(text))
    meta = _load_json(config.LECTURES_META_PATH, {})
    scores = []
    for lecture_id in sorted(urls):
        title_words = _norm_words(meta.get(lecture_id, {}).get("title", ""))
        if not title_words or not deck_words:
            continue
        # fraction of title words appearing in the deck's opening slides,
        # with fuzzy tolerance for OCR/diacritic drift
        hits = sum(
            1
            for w in title_words
            if w in deck_words or difflib.get_close_matches(w, deck_words, n=1, cutoff=0.85)
        )
        scores.append({"lecture_id": lecture_id, "score": round(hits / len(title_words), 3)})
    scores.sort(key=lambda s: -s["score"])
    return [s for s in scores if s["score"] >= 0.5][:3]
