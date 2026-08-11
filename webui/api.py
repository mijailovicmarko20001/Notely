"""JSON API for the Notely web UI."""

import asyncio
import json
import shutil
from pathlib import Path

from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse

from . import config, jobs, playlist, preflight, progress, review

router = APIRouter(prefix="/api")


def _load_json(path, default):
    try:
        with open(path) as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return default


# --- setup / settings -----------------------------------------------------

@router.get("/preflight")
def get_preflight():
    settings = config.read_settings()
    return preflight.run_preflight(
        ocr_lang=settings.get("OCR_LANG", "eng"),
        whisper_model=settings.get("WHISPER_MODEL", "medium"),
    )


@router.get("/settings")
def get_settings():
    return {"settings": config.read_settings(), "stage_defaults": config.DEFAULT_STAGE_OPTIONS}


@router.put("/settings")
def put_settings(body: dict):
    config.write_settings(body)
    return {"ok": True, "settings": config.read_settings()}


@router.post("/settings/test-key")
def test_key():
    key = config.get_api_key()
    if not key:
        raise HTTPException(400, "no API key saved")
    try:
        import anthropic

        client = anthropic.Anthropic(api_key=key, max_retries=1)
        client.messages.create(
            model="claude-haiku-4-5-20251001",
            max_tokens=8,
            messages=[{"role": "user", "content": "ping"}],
        )
        return {"ok": True}
    except Exception as e:
        return {"ok": False, "error": str(e)}


# --- lectures / sources ---------------------------------------------------

@router.get("/state")
def get_state():
    urls = _load_json(config.VIDEO_URLS_PATH, {})
    meta = _load_json(config.LECTURES_META_PATH, {})
    lectures = []
    for lecture_id in sorted(urls):
        stages = {
            s: progress.artifact_ok(s, lecture_id) for s in range(7)
        }
        deck = next(
            (p.name for ext in ("pptx", "pdf")
             for p in [config.SLIDES_DIR / f"{lecture_id}.{ext}"] if p.exists()),
            None,
        )
        lectures.append({
            "id": lecture_id,
            "url": urls[lecture_id],
            "title": meta.get(lecture_id, {}).get("title", lecture_id),
            "deck": deck,
            "stages": stages,
        })
    return {
        "lectures": lectures,
        "study_guide": progress.artifact_ok(7),
        "has_api_key": bool(config.get_api_key()),
        "job": jobs.MANAGER.snapshot(),
        "busy": jobs.MANAGER.busy,
    }


@router.post("/playlist/expand")
def expand(body: dict):
    url = (body.get("url") or "").strip()
    if not url:
        raise HTTPException(400, "url required")
    try:
        return {"entries": playlist.expand_playlist(url)}
    except playlist.PlaylistError as e:
        raise HTTPException(422, str(e))


@router.post("/lectures")
def set_lectures(body: dict):
    """body.entries: ordered [{title, url}] — assigns lecture01..NN."""
    entries = body.get("entries") or []
    if not entries:
        raise HTTPException(400, "entries required")
    urls, meta = {}, {}
    for i, e in enumerate(entries, start=1):
        lecture_id = f"lecture{i:02d}"
        urls[lecture_id] = e["url"]
        meta[lecture_id] = {"title": e.get("title", lecture_id)}
    config.INPUT_DIR.mkdir(parents=True, exist_ok=True)
    with open(config.VIDEO_URLS_PATH, "w") as f:
        json.dump(urls, f, indent=2)
    existing_meta = _load_json(config.LECTURES_META_PATH, {})
    for k, v in meta.items():
        existing_meta.setdefault(k, {}).update(v)
    with open(config.LECTURES_META_PATH, "w") as f:
        json.dump(existing_meta, f, ensure_ascii=False, indent=2)
    return {"ok": True, "lectures": list(urls)}


@router.post("/slides/upload-pool")
async def upload_pool(files: list[UploadFile] = File(...)):
    """'I don't know which deck covers which lecture' mode: save all decks to a
    pool, merge them into one combined PDF, and give every lecture that same
    combined deck. Stage 4's content matching then figures out per video which
    slides were actually shown — unshown slides are simply never matched
    (already the normal case, since decks can span lectures)."""
    import hashlib
    import re

    import fitz

    def norm_page_text(s: str) -> str:
        return re.sub(r"\s+", " ", (s or "").lower()).strip()

    urls = _load_json(config.VIDEO_URLS_PATH, {})
    if not urls:
        raise HTTPException(400, "add your lectures first — the combined deck is copied to each one")

    pool_dir = config.SLIDES_DIR / "_pool"
    pool_dir.mkdir(parents=True, exist_ok=True)
    for f in files:
        name = (f.filename or "").strip()
        if not name.lower().endswith(".pdf"):
            raise HTTPException(422, f"{name}: pool mode takes PDFs only — export PPTX decks to PDF first")
        content = await f.read()
        if not content:
            raise HTTPException(422, f"{name}: empty file")
        with open(pool_dir / Path(name).name, "wb") as out:
            out.write(content)

    # Merge every deck currently in the pool, in filename order, skipping
    # pages whose normalized text exactly matches one already kept. Course
    # decks repeat earlier material heavily (observed: 936 raw pages, only
    # 310 unique, on this project's real pool) — deduping here instead of
    # after stage 4's OCR makes the haystack stage 4 actually searches
    # ~3x smaller and less ambiguous, for free. Pages with no extractable
    # text (scanned/image-only slides) are never deduped against each other
    # — collapsing them on an empty-string hash match would wrongly merge
    # visually distinct slides, so they're always kept.
    pool_files = sorted(pool_dir.glob("*.pdf"))
    merged = fitz.open()
    seen_hashes = set()
    scanned_pages = 0
    for p in pool_files:
        with fitz.open(p) as src:
            for page_index in range(src.page_count):
                scanned_pages += 1
                text = norm_page_text(src[page_index].get_text())
                if text:
                    key = hashlib.md5(text.encode()).hexdigest()
                    if key in seen_hashes:
                        continue
                    seen_hashes.add(key)
                merged.insert_pdf(src, from_page=page_index, to_page=page_index)
    total_pages = merged.page_count
    merged_path = pool_dir / "_merged.pdf"
    merged.save(str(merged_path))
    merged.close()

    # same combined deck for every lecture (stages pair strictly by filename)
    for lecture_id in urls:
        shutil.copyfile(merged_path, config.SLIDES_DIR / f"{lecture_id}.pdf")
        (config.SLIDES_DIR / f"{lecture_id}.pptx").unlink(missing_ok=True)

    return {
        "ok": True,
        "pool_decks": [p.name for p in pool_files],
        "merged_pages": total_pages,
        "scanned_pages": scanned_pages,
        "duplicate_pages_skipped": scanned_pages - total_pages,
        "lectures": sorted(urls),
    }


@router.post("/slides/suggest")
async def suggest_slides(file: UploadFile = File(...)):
    """Suggest which lecture(s) a deck belongs to by comparing its first-slide
    text against the video titles. A pre-fill for the UI dropdown — never a
    silent decision."""
    import difflib
    import io
    import re

    def norm(s):
        return re.sub(r"[^\w\s]", " ", (s or "").lower()).split()

    content = await file.read()
    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    text = ""
    try:
        if ext == "pdf":
            import fitz

            with fitz.open(stream=content, filetype="pdf") as doc:
                # first two pages: title slide + first content slide
                text = " ".join(doc[i].get_text() for i in range(min(2, len(doc))))
        elif ext == "pptx":
            from pptx import Presentation

            prs = Presentation(io.BytesIO(content))
            for slide in list(prs.slides)[:2]:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        text += " " + shape.text_frame.text
    except Exception:
        pass  # unreadable deck -> no suggestions, UI falls back to order

    deck_words = set(norm(text))
    meta = _load_json(config.LECTURES_META_PATH, {})
    urls = _load_json(config.VIDEO_URLS_PATH, {})
    scores = []
    for lecture_id in sorted(urls):
        title_words = norm(meta.get(lecture_id, {}).get("title", ""))
        if not title_words or not deck_words:
            continue
        # fraction of title words appearing in the deck's opening slides,
        # with fuzzy tolerance for OCR/diacritic drift
        hits = sum(
            1 for w in title_words
            if w in deck_words or difflib.get_close_matches(w, deck_words, n=1, cutoff=0.85)
        )
        scores.append({"lecture_id": lecture_id, "score": round(hits / len(title_words), 3)})
    scores.sort(key=lambda s: -s["score"])
    return {"suggestions": [s for s in scores if s["score"] >= 0.5][:3]}


@router.post("/slides/upload")
async def upload_slides(file: UploadFile = File(...), lecture_ids: str = Form(...)):
    """One deck may cover several lectures: copy it under each lecture's name."""
    ids = [x.strip() for x in lecture_ids.split(",") if x.strip()]
    if not ids:
        raise HTTPException(400, "lecture_ids required")
    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    if ext not in ("pdf", "pptx"):
        raise HTTPException(422, "only .pdf and .pptx decks are supported")
    if ext == "pptx" and not preflight.check_soffice()["ok"]:
        raise HTTPException(422, "LibreOffice not available — convert this deck to PDF first")
    content = await file.read()
    if not content:
        raise HTTPException(422, "empty file")
    config.SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    saved = []
    for lecture_id in ids:
        dest = config.SLIDES_DIR / f"{lecture_id}.{ext}"
        with open(dest, "wb") as f:
            f.write(content)
        # remove a stale deck of the other extension so stage 02 pairs deterministically
        other = config.SLIDES_DIR / f"{lecture_id}.{'pdf' if ext == 'pptx' else 'pptx'}"
        other.unlink(missing_ok=True)
        saved.append(dest.name)
        meta = _load_json(config.LECTURES_META_PATH, {})
        meta.setdefault(lecture_id, {})["deck_original_name"] = file.filename
        with open(config.LECTURES_META_PATH, "w") as f:
            json.dump(meta, f, ensure_ascii=False, indent=2)
    return {"ok": True, "saved": saved}


# --- jobs -----------------------------------------------------------------

@router.post("/jobs")
def start_job(body: dict):
    if jobs.MANAGER.busy:
        raise HTTPException(409, "a job is already running")
    lecture_ids = body.get("lecture_ids") or []
    stages = sorted(set(body.get("stages") or []))
    if not stages or (not lecture_ids and stages != [7]):
        raise HTTPException(400, "lecture_ids and stages required")
    if any(s < 0 or s > 7 for s in stages):
        raise HTTPException(400, "stages must be 0-7")
    options = {**config.DEFAULT_STAGE_OPTIONS, **(body.get("options") or {})}
    tasks = jobs.build_tasks(
        lecture_ids, stages, options, bool(body.get("force")), bool(config.get_api_key())
    )
    if not tasks:
        raise HTTPException(400, "nothing to run (is stage 6 locked without an API key?)")
    job_id = jobs.MANAGER.start_job(tasks)
    return {"ok": True, "job_id": job_id}


@router.get("/jobs/current")
def current_job():
    return {"job": jobs.MANAGER.snapshot(), "busy": jobs.MANAGER.busy}


@router.post("/jobs/current/cancel")
def cancel_job():
    if not jobs.MANAGER.busy:
        raise HTTPException(409, "no job running")
    jobs.MANAGER.cancel()
    return {"ok": True}


@router.get("/jobs/current/events")
async def job_events(request: Request):
    last_id = request.headers.get("last-event-id")
    try:
        seq = int(last_id) if last_id else 0
    except ValueError:
        seq = 0

    async def gen(start_seq):
        seen = start_seq
        idle = 0.0
        while True:
            if await request.is_disconnected():
                return
            batch = jobs.MANAGER.events_since(seen)
            for evt in batch:
                seen = evt["seq"]
                yield f"id: {evt['seq']}\ndata: {json.dumps(evt)}\n\n"
            if batch:
                idle = 0.0
            else:
                idle += 0.25
                if idle >= 15.0:
                    yield ": heartbeat\n\n"
                    idle = 0.0
            await asyncio.sleep(0.25)

    return StreamingResponse(gen(seq), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache"})


# --- review ---------------------------------------------------------------

@router.get("/review/{lecture_id}")
def get_review(lecture_id: str):
    try:
        return review.get_review_data(lecture_id)
    except FileNotFoundError as e:
        raise HTTPException(404, str(e))


@router.post("/review/{lecture_id}")
def post_review(lecture_id: str, body: dict):
    corrections = body.get("corrections") or []
    if not corrections:
        raise HTTPException(400, "corrections required")
    if jobs.MANAGER.busy:
        raise HTTPException(409, "a job is already running — wait for it to finish")
    result = review.apply_corrections(lecture_id, corrections)
    # re-run downstream stages on the corrected timeline
    stages = [5, 6, 7] if config.get_api_key() else [5]
    tasks = jobs.build_tasks([lecture_id], stages, {}, True, bool(config.get_api_key()))
    job_id = jobs.MANAGER.start_job(tasks)
    return {"ok": True, **result, "job_id": job_id}


# --- guide ----------------------------------------------------------------

@router.get("/guide")
def get_guide():
    path = config.OUTPUT_DIR / "study_guide.md"
    if not path.exists():
        return {"exists": False}
    return {"exists": True, "markdown": path.read_text(), "download": "/files/study_guide.md"}


@router.get("/guide/pdf")
def guide_pdf():
    """Render the study guide to PDF (headless Chrome + MathJax) and return it."""
    import subprocess
    import sys as _sys

    from fastapi.responses import FileResponse

    md = config.OUTPUT_DIR / "study_guide.md"
    if not md.exists():
        raise HTTPException(404, "no study guide yet — run the pipeline first")
    pdf = config.OUTPUT_DIR / "study_guide.pdf"
    # regenerate when stale or missing
    if not pdf.exists() or pdf.stat().st_mtime < md.stat().st_mtime:
        script = config.SCRIPTS_DIR / "08_export_pdf.py"
        r = subprocess.run(
            [_sys.executable, str(script)], capture_output=True, text=True, timeout=300
        )
        if r.returncode != 0 or not pdf.exists():
            raise HTTPException(500, f"PDF export failed: {(r.stderr or r.stdout)[-500:]}")
    return FileResponse(pdf, media_type="application/pdf", filename="study_guide.pdf")
