"""Lecture-matching heuristic behind /slides/suggest (A2): "I don't know
which lecture this deck belongs to" mode -- suggest candidate lectures for
an uploaded deck by comparing its opening slides' text against each
lecture's video title. A pre-fill for the UI dropdown, never a silent
decision.

Moved out of webui/decks.py (Phase 6 of the cleanup plan): decks.py stays
focused on upload persistence and pool merge/dedup; this is a genuinely
separate concern (text-similarity matching, not file handling) that only
decks.py's own single upload endpoint used.
"""

import difflib
import logging
import re
import uuid
from pathlib import Path

from fastapi import UploadFile

from notely.io import load_json_or_default as _load_json
from . import config, decks

log = logging.getLogger("notely.lecture_match")


def _norm_words(s: str) -> list[str]:
    return re.sub(r"[^\w\s]", " ", (s or "").lower()).split()


def _extract_preview_text(path: Path, ext: str) -> str:
    """First couple pages/slides of text -- title slide + first content
    slide -- used as a fingerprint for lecture-matching. Unreadable decks
    log a warning and return "" so the UI just falls back to no
    suggestions, rather than the request failing outright."""
    text = ""
    try:
        if ext == "pdf":
            import pypdfium2 as pdfium

            with pdfium.PdfDocument(str(path)) as doc:
                parts = []
                for i in range(min(2, len(doc))):
                    page = doc[i]
                    textpage = page.get_textpage()
                    try:
                        parts.append(textpage.get_text_range())
                    finally:
                        textpage.close()
                        page.close()
                text = " ".join(parts)
        elif ext == "pptx":
            from pptx import Presentation

            prs = Presentation(str(path))
            parts = []
            for slide in list(prs.slides)[:2]:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            text = " ".join(parts)
    except Exception:
        log.warning("could not extract preview text from .%s upload", ext, exc_info=True)
    return text


async def suggest_lectures_for_upload(
    file: UploadFile, max_bytes: int = config.MAX_UPLOAD_BYTES
) -> list[dict]:
    """Suggest which lecture(s) a deck belongs to by comparing its opening
    slides' text against the video titles. A pre-fill for the UI dropdown,
    never a silent decision."""
    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    urls = config.load_video_urls()
    if ext not in ("pdf", "pptx") or not urls:
        return []

    tmp = config.SLIDES_DIR / f"_suggest_{uuid.uuid4().hex}.{ext}"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    try:
        await decks.save_upload_stream(file, tmp, max_bytes, allow_empty=True)
        text = _extract_preview_text(tmp, ext)
    finally:
        tmp.unlink(missing_ok=True)

    deck_words = set(_norm_words(text))
    meta = _load_json(config.LECTURES_META_PATH, {})
    scores = []
    for lecture_id in sorted(urls):
        title_words = _norm_words(meta.get(lecture_id, {}).get("title", ""))
        if not title_words or not deck_words:
            continue
        # fraction of title words appearing in the deck's opening slides,
        # with fuzzy tolerance for OCR/diacritic drift
        hits = sum(
            1
            for w in title_words
            if w in deck_words or difflib.get_close_matches(w, deck_words, n=1, cutoff=0.85)
        )
        scores.append({"lecture_id": lecture_id, "score": round(hits / len(title_words), 3)})
    scores.sort(key=lambda s: -s["score"])
    return [s for s in scores if s["score"] >= 0.5][:3]
